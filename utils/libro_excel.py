"""
Reemplaza el botón "Actualizar" de Power Query, que solo funcionaba en Excel
sobre Windows. Vuelca las tablas de la base SQLite en las hojas ocultas del
libro; las hojas visibles, que son fórmulas sobre esas tablas, se recalculan
solas al abrirlo.

"""
import re
import shutil
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

LIBRO_POR_DEFECTO = Path.home() / "Documentos" / "CharcuteriaHYE.xlsx"

# Hoja oculta → columnas que se traen de la base, en el orden en que están
# en el libro. Lo que venga después de estas columnas es manual y no se toca.
HOJAS = {
    "sales":               ["total", "payment_method", "created_at", "status"],
    "expenses":            ["amount", "date"],
    "losses":              ["product_id", "quantity", "unit_cost", "total_cost", "reason", "created_at"],
    "purchases":           ["total", "iva", "shipping", "date", "lote_id"],
    "client_transactions": ["transaction_type", "amount", "description", "sale_id", "created_at"],
    "products":            ["id", "name", "price", "stock", "cost_price"],
    "sale_details":        ["id", "sale_id", "product_id", "quantity", "unit_price",
                            "sale_price", "subtotal", "cost_price"],
}

COLUMNAS_FECHA = {"created_at", "date"}
FORMATO_FECHA  = "m/d/yy h:mm"

# Hojas con columnas escritas a mano que hay que conservar entre actualizaciones.
# La clave es la columna por la que se emparejan las filas.
COLUMNAS_MANUALES = {"products": "id"}

# "Pendientes" tampoco es una hoja escrita a mano: era una consulta SQL con
# columnas calculadas. Se reproduce igual que la tenía Power Query.
CONSULTA_PENDIENTES = """
SELECT c.name,
       c.total_debt,
       (SELECT DATE(s.created_at)
          FROM sales s
         WHERE s.client_id = c.id AND s.status = 'pending'
         ORDER BY s.created_at ASC
         LIMIT 1) AS debe_desde
  FROM clients c
 WHERE c.total_debt > 0
 ORDER BY c.total_debt DESC
"""

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio",
         "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

FORMATO_MONEDA = '_-"$"* #,##0_-;\\-"$"* #,##0_-;_-"$"* "-"??_-;_-@_-'
FORMATO_DIA    = "mm-dd-yy"


def _actualizar_pendientes(libro, conexion) -> int:
    """Lista de deudores, con los días que llevan debiendo calculados a hoy."""
    hoja = libro["Pendientes"]
    filas = conexion.execute(CONSULTA_PENDIENTES).fetchall()
    hoy = datetime.now().date()
    anteriores = hoja.max_row

    for i, (nombre, deuda, desde) in enumerate(filas, start=2):
        fecha = _a_fecha(desde)
        hoja.cell(i, 1).value = nombre
        hoja.cell(i, 2).value = deuda
        hoja.cell(i, 2).number_format = FORMATO_MONEDA
        if isinstance(fecha, datetime):
            hoja.cell(i, 3).value = f"{fecha.day} de {MESES[fecha.month - 1]}"
            hoja.cell(i, 4).value = (hoy - fecha.date()).days
            hoja.cell(i, 5).value = fecha
            hoja.cell(i, 5).number_format = FORMATO_DIA
        else:
            for col in (3, 4, 5):
                hoja.cell(i, col).value = None

    for fila in range(len(filas) + 2, anteriores + 1):
        for col in range(1, 6):
            hoja.cell(fila, col).value = None

    if "clients" in hoja.tables:
        hoja.tables["clients"].ref = f"A1:E{len(filas) + 1}"

    # Fila de total, siempre justo debajo de los datos. Antes estaba escrita a
    # mano en una fila fija y quedaba descolgada cada vez que cambiaba el
    # número de deudores.
    fila_total = len(filas) + 2
    hoja.cell(fila_total, 1).value = "Total"
    hoja.cell(fila_total, 2).value = f"=SUM(B2:B{len(filas) + 1})"
    hoja.cell(fila_total, 2).number_format = FORMATO_MONEDA

    return len(filas)


def _a_fecha(valor):
    """SQLite guarda las fechas como texto; el libro las necesita como fechas
    de verdad, porque las fórmulas hacen comparaciones y restas con ellas."""
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor
    texto = str(valor).strip()
    for formato in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto, formato)
        except ValueError:
            continue
    return valor  # se queda como texto si tiene un formato inesperado


def _traducir_xlookup(libro) -> int:
    """XLOOKUP no existe en LibreOffice 24.2 (da #¿NOMBRE?). INDEX/MATCH hace
    lo mismo y funciona igual en Excel."""
    patron = re.compile(
        r"_xlfn\.XLOOKUP\(([^,]+),(\w+\[[^\]]+\]),(\w+\[[^\]]+\])\)"
    )
    cambiadas = 0
    for hoja in libro.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                if isinstance(celda.value, str) and "_xlfn.XLOOKUP" in celda.value:
                    nueva = patron.sub(r"INDEX(\3,MATCH(\1,\2,0))", celda.value)
                    if nueva != celda.value:
                        celda.value = nueva
                        cambiadas += 1
    return cambiadas


def actualizar(ruta_libro: Path, ruta_bd: Path) -> None:
    if not ruta_libro.exists():
        raise FileNotFoundError(f"No encuentro el libro: {ruta_libro}")
    if not ruta_bd.exists():
        raise FileNotFoundError(f"No encuentro la base de datos: {ruta_bd}")

    respaldo = ruta_libro.with_suffix(".xlsx.anterior")
    shutil.copy2(ruta_libro, respaldo)

    conexion = sqlite3.connect(ruta_bd)
    libro = openpyxl.load_workbook(ruta_libro)

    for nombre, columnas in HOJAS.items():
        if nombre not in libro.sheetnames:
            print(f"  · {nombre}: no está en el libro, la salto")
            continue
        hoja = libro[nombre]

        # Guardar las columnas manuales antes de reescribir
        manuales = {}
        clave = COLUMNAS_MANUALES.get(nombre)
        if clave:
            i_clave = columnas.index(clave)
            for fila in range(2, hoja.max_row + 1):
                id_fila = hoja.cell(fila, i_clave + 1).value
                if id_fila is not None:
                    manuales[id_fila] = [
                        hoja.cell(fila, c).value
                        for c in range(len(columnas) + 1, hoja.max_column + 1)
                    ]

        filas = conexion.execute(
            f"SELECT {', '.join(columnas)} FROM {nombre} ORDER BY id"
        ).fetchall()

        anteriores = hoja.max_row
        for i, registro in enumerate(filas, start=2):
            for j, (col, valor) in enumerate(zip(columnas, registro), start=1):
                celda = hoja.cell(i, j)
                if col in COLUMNAS_FECHA:
                    celda.value = _a_fecha(valor)
                    celda.number_format = FORMATO_FECHA
                else:
                    celda.value = valor
            # Devolver a su sitio lo escrito a mano
            if clave:
                for k, valor in enumerate(manuales.get(registro[i_clave], []),
                                          start=len(columnas) + 1):
                    hoja.cell(i, k).value = valor

        # Borrar las filas sobrantes si la base tiene menos registros que antes
        for fila in range(len(filas) + 2, anteriores + 1):
            for col in range(1, hoja.max_column + 1):
                hoja.cell(fila, col).value = None

        # Ajustar el rango de la tabla, o las referencias tipo tabla[columna]
        # dejarían fuera las filas nuevas
        if nombre in hoja.tables:
            ultima_col = hoja.tables[nombre].ref.split(":")[1].rstrip("0123456789")
            hoja.tables[nombre].ref = f"A1:{ultima_col}{len(filas) + 1}"

        print(f"  · {nombre}: {len(filas)} filas")

    print(f"  · Pendientes: {_actualizar_pendientes(libro, conexion)} deudores")

    cambiadas = _traducir_xlookup(libro)
    if cambiadas:
        print(f"  · {cambiadas} fórmulas XLOOKUP traducidas a INDEX/MATCH")

    libro.save(ruta_libro)
    conexion.close()
    print(f"\nListo: {ruta_libro}")
    print(f"Copia anterior: {respaldo}")


