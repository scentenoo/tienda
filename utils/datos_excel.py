"""
Genera datos.xlsx: un libro con una tabla por consulta, para que Excel Online
lo lea desde OneDrive con el conector "Libro de Excel".

Las columnas son exactamente las que las consultas producían con ODBC, para
que las fórmulas de las hojas visibles sigan funcionando igual.
"""
import sqlite3
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

CONSULTAS = {
    "sales":               "SELECT total, payment_method, created_at, status FROM sales ORDER BY id",
    "expenses":            "SELECT amount, date FROM expenses ORDER BY id",
    "losses":              "SELECT product_id, quantity, unit_cost, total_cost, reason, created_at FROM losses ORDER BY id",
    "purchases":           "SELECT total, iva, shipping, date, lote_id FROM purchases ORDER BY id",
    "client_transactions": "SELECT transaction_type, amount, description, sale_id, created_at FROM client_transactions ORDER BY id",
    "products":            "SELECT id, name, price, stock, cost_price FROM products ORDER BY id",
    "sale_details":        "SELECT id, sale_id, product_id, quantity, unit_price, sale_price, subtotal, cost_price FROM sale_details ORDER BY id",
}

COLUMNAS_FECHA = {"created_at", "date"}
FORMATO_FECHA  = "m/d/yy h:mm"

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio",
         "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]

# La consulta "clients" alimentaba la hoja Pendientes y traía columnas
# calculadas que Power Query armaba aparte.
SQL_CLIENTS = """
SELECT c.name, c.total_debt,
       (SELECT DATE(s.created_at) FROM sales s
         WHERE s.client_id = c.id AND s.status = 'pending'
         ORDER BY s.created_at ASC LIMIT 1) AS debe_desde
  FROM clients c WHERE c.total_debt > 0 ORDER BY c.total_debt DESC
"""


def _a_fecha(valor):
    if valor in (None, ""):
        return None
    for formato in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(valor).strip(), formato)
        except ValueError:
            continue
    return valor


def _escribir(hoja, encabezados, filas):
    hoja.append(encabezados)
    for registro in filas:
        hoja.append(list(registro))
    for i, nombre in enumerate(encabezados, start=1):
        if nombre in COLUMNAS_FECHA:
            for fila in range(2, len(filas) + 2):
                hoja.cell(fila, i).number_format = FORMATO_FECHA
    ultima = openpyxl.utils.get_column_letter(len(encabezados))
    tabla = Table(displayName=hoja.title, ref=f"A1:{ultima}{len(filas) + 1}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    hoja.add_table(tabla)


def generar(salida: Path, bd: Path) -> int:
    """Devuelve el número de tablas escritas."""
    conexion = sqlite3.connect(bd)
    libro = openpyxl.Workbook()
    libro.remove(libro.active)

    for nombre, sql in CONSULTAS.items():
        cursor = conexion.execute(sql)
        encabezados = [d[0] for d in cursor.description]
        filas = [
            [_a_fecha(v) if c in COLUMNAS_FECHA else v for c, v in zip(encabezados, r)]
            for r in cursor.fetchall()
        ]
        _escribir(libro.create_sheet(nombre), encabezados, filas)

    hoy = datetime.now().date()
    filas = []
    for nombre_cliente, deuda, desde in conexion.execute(SQL_CLIENTS):
        fecha = _a_fecha(desde)
        if isinstance(fecha, datetime):
            filas.append([nombre_cliente, deuda,
                          f"{fecha.day} de {MESES[fecha.month - 1]}",
                          (hoy - fecha.date()).days, fecha])
        else:
            filas.append([nombre_cliente, deuda, None, None, None])
    hoja = libro.create_sheet("clients")
    _escribir(hoja, ["Nombre", "Deuda_Total", "DebeDesde", "Dias en Deuda", "Debe_Desde"], filas)
    for fila in range(2, len(filas) + 2):
        hoja.cell(fila, 5).number_format = "mm-dd-yy"

    libro.save(salida)
    conexion.close()
    return len(CONSULTAS) + 1
